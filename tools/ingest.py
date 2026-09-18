#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Ingest Mike DuBois' weekly picks workbook into data/<season>/week-NN.json.

    uv run tools/ingest.py ~/Downloads/week3.xlsx --season 2026 --week 3
    python3 tools/ingest.py week3.xlsx --season 2026 --week 3     # if openpyxl is installed

Workbook shape (sheet "Picks"): row 1 is a header of column pairs —
`<participant name>`, `Points` — one pair per entry. Rows below hold one game
per row: the full name of the team that entry picked and its confidence points
(1..N, each used once). A blank row, then a `MNF Pts.` row with the tiebreaker
(predicted total points in the Monday night game). The sheet does not say who
each team is playing, so games are matched to ESPN's schedule by team.

Bad rows warn, they do not crash: an entry with a duplicate or missing
confidence number is kept and flagged in `warnings` so the site can show it.

Also refreshes data/<season>/results/week-NN.json (the same scoreboard fetch
gives us kickoffs and lines) and the week indexes the site reads.
"""
import argparse
import json
import os
import re
import sys
from collections import Counter
from datetime import datetime, timezone

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from teams import FULL_NAME, normalize  # noqa: E402
import scores  # noqa: E402

ROOT = scores.ROOT
DATA = scores.DATA


def clean_name(s):
    s = str(s or "").replace("﻿", "")
    return re.sub(r"\s+", " ", s).strip()


def _is_points_header(v):
    return clean_name(v).lower() in ("points", "pts", "point")


def read_workbook(path):
    """-> (names, game_rows, mnf_row) straight from the sheet, un-normalised.

    names:     [participant, ...] in column order
    game_rows: [[(team_text, points), ...per participant], ...per game row]
    mnf_row:   [tiebreaker value, ...per participant] or None
    """
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["Picks"] if "Picks" in wb.sheetnames else wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if not rows:
        raise SystemExit(f"{path}: empty sheet")

    header = rows[0]
    cols = []   # (name_col, points_col) per participant
    c = 0
    while c < len(header):
        name = clean_name(header[c])
        if name and not _is_points_header(name):
            pts_col = c + 1 if c + 1 < len(header) and _is_points_header(header[c + 1]) else None
            cols.append((name, c, pts_col))
            c += 2 if pts_col is not None else 1
        else:
            c += 1
    if not cols:
        raise SystemExit(f"{path}: no participant columns found in row 1")

    game_rows, mnf_row = [], None
    for r in rows[1:]:
        texts = [clean_name(r[nc]) for _, nc, _ in cols if nc < len(r)]
        if not any(texts):
            continue
        if any(re.search(r"mnf|tie ?break", t, re.I) for t in texts):
            mnf_row = [r[pc] if pc is not None and pc < len(r) else None for _, _, pc in cols]
            continue
        if any(normalize(t) for t in texts):
            game_rows.append([
                (r[nc] if nc < len(r) else None, r[pc] if pc is not None and pc < len(r) else None)
                for _, nc, pc in cols
            ])
    names = [n for n, _, _ in cols]
    return names, game_rows, mnf_row


def _num(v):
    if v is None or v == "":
        return None
    try:
        f = float(str(v).strip())
    except ValueError:
        return None
    return int(f) if f.is_integer() else f


def build_entries(names, game_rows, mnf_row):
    n_games = len(game_rows)
    entries, seen = [], Counter()
    for i, name in enumerate(names):
        warnings = []
        picks, points = [], []
        for g, row in enumerate(game_rows):
            team_text, pts = row[i]
            abbr = normalize(team_text)
            if abbr is None:
                warnings.append(f"game {g + 1}: unrecognised pick {team_text!r}")
            p = _num(pts)
            if not isinstance(p, int) or not 1 <= p <= n_games:
                warnings.append(f"game {g + 1}: bad points {pts!r}")
                p = None
            picks.append(abbr)
            points.append(p)
        used = [p for p in points if p is not None]
        if sorted(used) != list(range(1, n_games + 1)):
            dupes = sorted(p for p, k in Counter(used).items() if k > 1)
            missing = sorted(set(range(1, n_games + 1)) - set(used))
            warnings.append(f"points not 1..{n_games} once each"
                            + (f"; duplicates {dupes}" if dupes else "")
                            + (f"; missing {missing}" if missing else ""))
        mnf = _num(mnf_row[i]) if mnf_row is not None else None
        if mnf is None:
            warnings.append("no MNF tiebreaker")

        seen[name] += 1
        display = name if seen[name] == 1 else f"{name} ({seen[name]})"
        if seen[name] > 1:
            warnings.append(f"duplicate participant name {name!r}")
        entries.append({
            "name": display,
            "picks": picks,
            "points": points,
            "mnf_total": mnf,
            "warnings": warnings,
        })
    return entries


def build_games(game_rows, espn_games):
    """One record per sheet row, matched to an ESPN event by the teams picked."""
    by_id = {g["espn_id"]: g for g in espn_games}
    used = set()
    games = []
    for g, row in enumerate(game_rows):
        teams = Counter(normalize(t) for t, _ in row if normalize(t))
        picked = set(teams)
        # best overlap wins; an event no other row has claimed breaks ties
        match, best = None, (0, False)
        for ev in espn_games:
            score = (len(picked & {ev["away"], ev["home"]}), ev["espn_id"] not in used)
            if score > best:
                match, best = ev, score
        best = best[0]
        rec = {"idx": g, "picked_teams": sorted(picked)}
        if match and best:
            used.add(match["espn_id"])
            rec.update({
                "espn_id": match["espn_id"],
                "away": match["away"], "home": match["home"],
                "away_name": match["away_name"], "home_name": match["home_name"],
                "kickoff": match["kickoff"],
            })
            extra = picked - {match["away"], match["home"]}
            if extra:
                rec["warning"] = f"picks for {sorted(extra)} do not belong to {match['short_name']}"
        else:
            # no schedule available: keep the teams we saw so the row still renders
            ordered = sorted(picked)
            rec.update({
                "espn_id": None,
                "away": ordered[0] if ordered else None,
                "home": ordered[1] if len(ordered) > 1 else None,
                "away_name": FULL_NAME.get(ordered[0]) if ordered else None,
                "home_name": FULL_NAME.get(ordered[1]) if len(ordered) > 1 else None,
                "kickoff": None,
                "warning": "not matched to an ESPN game",
            })
        games.append(rec)
    return games


def update_indexes(season, week):
    sdir = os.path.join(DATA, str(season))
    os.makedirs(sdir, exist_ok=True)
    sidx = os.path.join(sdir, "index.json")
    weeks = set()
    if os.path.exists(sidx):
        with open(sidx) as f:
            weeks = set(int(w) for w in json.load(f).get("weeks", []))
    weeks.add(int(week))
    with open(sidx, "w") as f:
        json.dump({"season": int(season), "weeks": sorted(weeks)}, f, indent=1)
        f.write("\n")

    ridx = os.path.join(DATA, "index.json")
    seasons = {}
    if os.path.exists(ridx):
        with open(ridx) as f:
            seasons = {int(k): v for k, v in json.load(f).get("seasons", {}).items()}
    seasons[int(season)] = sorted(weeks)
    latest_season = max(seasons)
    with open(ridx, "w") as f:
        json.dump({
            "seasons": {str(k): v for k, v in sorted(seasons.items())},
            "latest": {"season": latest_season, "week": max(seasons[latest_season])},
        }, f, indent=1)
        f.write("\n")


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("xlsx")
    ap.add_argument("--season", type=int, required=True)
    ap.add_argument("--week", type=int, required=True)
    ap.add_argument("--no-espn", action="store_true", help="skip the ESPN schedule fetch")
    ap.add_argument("--espn-file", help="use a saved scoreboard JSON instead of fetching")
    args = ap.parse_args()

    names, game_rows, mnf_row = read_workbook(args.xlsx)
    entries = build_entries(names, game_rows, mnf_row)

    espn_games = []
    if args.espn_file:
        with open(args.espn_file) as f:
            d = json.load(f)
        events = d["content"]["sbData"]["events"] if "content" in d else d
        espn_games = scores.parse_events(events)
    elif not args.no_espn:
        try:
            espn_games = scores.parse_events(scores.fetch_scoreboard(args.season, args.week))
        except Exception as e:  # noqa: BLE001 — degrade, don't crash
            print(f"warning: ESPN fetch failed ({e}); games left unmatched", file=sys.stderr)

    games = build_games(game_rows, espn_games)

    doc = {
        "season": args.season,
        "week": args.week,
        "source": os.path.basename(args.xlsx),
        "ingested_at": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "n_games": len(games),
        "games": games,
        "entries": entries,
    }
    out = os.path.join(DATA, str(args.season), f"week-{args.week:02d}.json")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, "w") as f:
        json.dump(doc, f, indent=1)
        f.write("\n")

    if espn_games:
        scores.write_results(args.season, args.week, espn_games)
    update_indexes(args.season, args.week)

    # ---- report ----
    bad = [e for e in entries if e["warnings"]]
    unmatched = [g for g in games if not g.get("espn_id")]
    print(f"{os.path.relpath(out, ROOT)}: {len(entries)} entries, {len(games)} games"
          f" ({len(games) - len(unmatched)} matched to ESPN), {len(bad)} entries with warnings")
    for g in games:
        if g.get("warning"):
            print(f"  game {g['idx'] + 1} {g['picked_teams']}: {g['warning']}")
    for e in bad:
        print(f"  {e['name']}: " + "; ".join(e["warnings"]))


if __name__ == "__main__":
    main()
